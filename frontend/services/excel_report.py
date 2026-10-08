"""
excel_report.py — Erzeugt die Ergebnis-Excel-Datei einer Evaluation.

Unabhängig von ragas/LangChain testbar: Eingabe ist eine Liste von
dict-Zeilen (ROW_SCHEMA), Ausgabe eine .xlsx-Datei mit den Blättern
"Evaluation Results", "Summary", "Metric Overview" (mit Balkendiagramm),
"Issues" und "Beantwortbarkeit".

ROW_SCHEMA (ein dict pro Testfall):
    nr, produktbereich, produkt, artikelnummer, schwierigkeit, fragetyp,
    frage, referenzantwort, generierte_antwort,
    kontext (str, '\\n---\\n'-getrennt), kontext_anzahl (int),
    quellen (str, kurze Zusammenfassung),
    query_class, latency_ms, cache_hit, system_faithfulness_score,
    verified, fallback_reason, nicht_beantwortbar, data_quality_flags (str),
    ragas_faithfulness, ragas_answer_relevancy, ragas_context_precision,
    ragas_context_recall,
    faithfulness_hinweis, answer_relevancy_hinweis,
    context_precision_hinweis, context_recall_hinweis,
    fehler (str) — technischer Fehler bei Abruf oder Berechnung
"""
from __future__ import annotations

import re
import statistics
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

METRICS = [
    ("ragas_faithfulness", "Faithfulness"),
    ("ragas_answer_relevancy", "Answer Relevancy"),
    ("ragas_context_precision", "Context Precision"),
    ("ragas_context_recall", "Context Recall"),
]
# Zielwerte je Metrik und Ausbaustufe (Kap. 4.7.2). Gelten nur für
# beantwortbare Testfälle: Eine korrekte Ablehnung (F-13) hat regulär eine
# niedrige Answer Relevancy.
ARCHITEKTUR_ZIELE = {
    "ragas_faithfulness": {"Stufe 1 (Phase 2b)": 0.78, "Stufe 2 (Phase 3a, NF-04)": 0.85},
    "ragas_context_precision": {"Stufe 1 (Phase 2b)": 0.74, "Stufe 2 (Phase 3a)": 0.80},
    "ragas_answer_relevancy": {"Stufe 2 (Phase 3a)": 0.80},
}

# Erkennung von Ablehnungsantworten, identisch zu services/evaluation.py.
# Bewusst dupliziert, damit dieses Modul ohne ragas-Abhängigkeit importierbar bleibt.
_VERWEIGERUNGS_MARKER = (
    "nicht belegt", "nicht verfügbar", "keine information", "keine relevanten dokumente",
    "nicht im kontext", "ist nicht bekannt", "kann ich nicht beantworten",
    "liegen keine daten vor", "wurde nicht gefunden",
)


HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ISSUE_FILL = PatternFill(start_color="FCE4E4", end_color="FCE4E4", fill_type="solid")

# Steuerzeichen, die openpyxl ablehnt (IllegalCharacterError); Tab, LF und CR
# sind erlaubt. Treten bei fehlerhaft extrahiertem PDF-/OCR-Text auf.
_ILLEGAL_EXCEL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

def _antwort_verweigert(antwort: str) -> bool:
    text = (antwort or "").lower()
    return any(marker in text for marker in _VERWEIGERUNGS_MARKER)

def _sanitize_value(value: Any) -> Any:
    """Entfernt in Excel-Zellen unzulässige Steuerzeichen."""
    if isinstance(value, str):
        return _ILLEGAL_EXCEL_CHARS_RE.sub("", value)
    return value


def _sanitize_rows(rows: list[dict]) -> list[dict]:
    return [
        {k: _sanitize_value(v) for k, v in row.items()}
        for row in rows
    ]


def _style_header(ws: Worksheet, ncols: int, row: int = 1) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _autosize(ws: Worksheet, df: pd.DataFrame, max_width: int = 60) -> None:
    for i, col in enumerate(df.columns, start=1):
        lengths = [len(str(v)) for v in df[col].tolist()] or [0]
        width = min(max_width, max(10, max(lengths), len(str(col)) + 2))
        ws.column_dimensions[get_column_letter(i)].width = width


def _metric_series(rows: list[dict], key: str) -> list[float]:
    return [r[key] for r in rows if isinstance(r.get(key), (int, float))]


def _agg(values: list[float]) -> dict[str, Any]:
    if not values:
        return {"n": 0, "mean": None, "min": None, "max": None, "std": None}
    return {
        "n": len(values),
        "mean": round(statistics.fmean(values), 4),
        "min": round(min(values), 4),
        "max": round(max(values), 4),
        "std": round(statistics.pstdev(values), 4) if len(values) > 1 else 0.0,
    }


def build_excel(
    rows: list[dict],
    output_path: str | Path,
    run_meta: dict[str, Any] | None = None,
) -> Path:
    """Schreibt die vollständige .xlsx-Datei.

    Args:
        rows: Testfallzeilen gemäß ROW_SCHEMA.
        output_path: Zielpfad; das Verzeichnis wird bei Bedarf angelegt.
        run_meta: Optionale Laufmetadaten, oben im Blatt "Summary" eingefügt.

    Returns:
        Pfad der geschriebenen Datei.
    """
    rows = _sanitize_rows(rows)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    run_meta = run_meta or {}
    if run_meta:
        run_meta = {k: _sanitize_value(v) for k, v in run_meta.items()}

    # Blatt "Evaluation Results"
    results_cols = [
        "nr", "produktbereich", "produkt", "artikelnummer", "schwierigkeit",
        "fragetyp", "nicht_beantwortbar", "frage", "referenzantwort",
        "generierte_antwort", "kontext_anzahl", "kontext", "quellen",
        "query_class", "latency_ms", "cache_hit", "latency_ms_warm", "cache_hit_warm",
        "system_faithfulness_score",
        "verified", "fallback_reason",
        "ragas_faithfulness", "faithfulness_delta", "ragas_answer_relevancy",
        "ragas_context_precision", "ragas_context_recall",
        "faithfulness_hinweis", "answer_relevancy_hinweis",
        "context_precision_hinweis", "context_recall_hinweis",
        "data_quality_flags", "fehler",
    ]
    df_results = pd.DataFrame(rows)
    for c in results_cols:
        if c not in df_results.columns:
            df_results[c] = None
    df_results = df_results[results_cols].sort_values("nr")

    # Blatt "Summary"
    n_total = len(rows)
    n_errors = sum(1 for r in rows if r.get("fehler"))
    n_negativ = sum(1 for r in rows if r.get("nicht_beantwortbar"))
    n_flagged = sum(1 for r in rows if r.get("data_quality_flags"))

    summary_rows = [
        {"Kennzahl": "Anzahl Testfaelle gesamt", "Wert": n_total},
        {"Kennzahl": "Davon technische Fehler beim Abruf", "Wert": n_errors},
        {"Kennzahl": "Davon Negativfaelle (nicht beantwortbar)", "Wert": n_negativ},
        {"Kennzahl": "Davon mit Datenqualitaets-Hinweis", "Wert": n_flagged},
    ]
    for key, label in METRICS:
        agg = _agg(_metric_series(rows, key))
        summary_rows.append({
            "Kennzahl": f"{label} — n bewertet", "Wert": agg["n"],
        })
        summary_rows.append({
            "Kennzahl": f"{label} — Mittelwert", "Wert": agg["mean"],
        })
        summary_rows.append({
            "Kennzahl": f"{label} — Min", "Wert": agg["min"],
        })
        summary_rows.append({
            "Kennzahl": f"{label} — Max", "Wert": agg["max"],
        })
        summary_rows.append({
            "Kennzahl": f"{label} — Std.-Abw.", "Wert": agg["std"],
        })

    # Latenz: Standardmessung und, falls vorhanden, Warm-Cache-Messung
    # (latency_ms_warm ist nur bei aktivem Latenzvergleich gesetzt).
    lat_agg = _agg(_metric_series(rows, "latency_ms"))
    summary_rows.append({"Kennzahl": "Latenz (Standard) — Mittelwert ms", "Wert": lat_agg["mean"]})
    summary_rows.append({"Kennzahl": "Latenz (Standard) — Min ms", "Wert": lat_agg["min"]})
    summary_rows.append({"Kennzahl": "Latenz (Standard) — Max ms", "Wert": lat_agg["max"]})

    warm_agg = _agg(_metric_series(rows, "latency_ms_warm"))
    if warm_agg["n"]:
        n_hits = sum(1 for r in rows if r.get("cache_hit_warm"))
        summary_rows.append({"Kennzahl": "Latenz warm (Cache) — n gemessen", "Wert": warm_agg["n"]})
        summary_rows.append({"Kennzahl": "Latenz warm (Cache) — Mittelwert ms", "Wert": warm_agg["mean"]})
        summary_rows.append({"Kennzahl": "Latenz warm (Cache) — Min ms", "Wert": warm_agg["min"]})
        summary_rows.append({"Kennzahl": "Latenz warm (Cache) — Max ms", "Wert": warm_agg["max"]})
        summary_rows.append({
            "Kennzahl": "Warm-Cache-Trefferquote %",
            "Wert": round(100 * n_hits / warm_agg["n"], 1),
        })
        if lat_agg["mean"] and warm_agg["mean"]:
            summary_rows.append({
                "Kennzahl": "Speedup-Faktor (kalt/warm)",
                "Wert": round(lat_agg["mean"] / warm_agg["mean"], 2),
            })

    df_summary = pd.DataFrame(summary_rows)

    # Blatt "Metric Overview"
    overview_rows = []
    for key, label in METRICS:
        agg = _agg(_metric_series(rows, key))
        overview_rows.append({
            "Metrik": label,
            "Durchschnitt": agg["mean"],
            "Anzahl bewertet": agg["n"],
            "Anzahl N/A": n_total - agg["n"],
            "Min": agg["min"],
            "Max": agg["max"],
            "Std.-Abw.": agg["std"],
        })
    df_overview = pd.DataFrame(overview_rows)

    # Blatt "Issues"
    issue_rows = []
    for r in rows:
        problems = []
        if r.get("fehler"):
            problems.append(f"Technischer Fehler: {r['fehler']}")
        for hint_key, label in [
            ("faithfulness_hinweis", "Faithfulness"),
            ("answer_relevancy_hinweis", "Answer Relevancy"),
            ("context_precision_hinweis", "Context Precision"),
            ("context_recall_hinweis", "Context Recall"),
        ]:
            if r.get(hint_key):
                problems.append(f"{label} = N/A: {r[hint_key]}")
        if r.get("data_quality_flags"):
            problems.append(f"Testfall-Datenqualitaet: {r['data_quality_flags']}")
        if not r.get("kontext_anzahl"):
            problems.append("Kein Kontext von der API zurueckgegeben.")

        if problems:
            issue_rows.append({
                "nr": r.get("nr"),
                "frage": r.get("frage"),
                "nicht_beantwortbar": r.get("nicht_beantwortbar"),
                "Auffaelligkeiten": " | ".join(problems),
            })
    df_issues = pd.DataFrame(issue_rows) if issue_rows else pd.DataFrame(
        columns=["nr", "frage", "nicht_beantwortbar", "Auffaelligkeiten"]
    )
    # Blatt "Beantwortbarkeit" (F-13, Kap. 4.7.2): Zielwertprüfung nur über
    # beantwortbare Fälle; Ablehnungsquote über nicht beantwortbare Fälle.
    beantwortbar_rows = [r for r in rows if not r.get("nicht_beantwortbar")]
    nicht_beantwortbar_rows = [r for r in rows if r.get("nicht_beantwortbar")]

    ziel_rows = []
    for key, label in METRICS:
        agg = _agg(_metric_series(beantwortbar_rows, key))
        zeile = {
            "Metrik": label,
            "Durchschnitt (nur beantwortbare Faelle)": agg["mean"],
            "n bewertet": agg["n"],
        }
        ziele = ARCHITEKTUR_ZIELE.get(key, {})
        for stufe, ziel_wert in ziele.items():
            zeile[f"Ziel {stufe}"] = ziel_wert
            zeile[f"Erreicht {stufe}"] = (
                "Ja" if isinstance(agg["mean"], (int, float)) and agg["mean"] >= ziel_wert else "Nein"
            )
        ziel_rows.append(zeile)
    df_ziele = pd.DataFrame(ziel_rows)

    n_neg = len(nicht_beantwortbar_rows)
    n_korrekt = sum(1 for r in nicht_beantwortbar_rows if _antwort_verweigert(r.get("generierte_antwort", "")))
    halluziniert_nr = [
        r.get("nr") for r in nicht_beantwortbar_rows
        if not _antwort_verweigert(r.get("generierte_antwort", ""))
    ]
    df_ablehnung = pd.DataFrame([{
        "Anzahl nicht-beantwortbare Faelle": n_neg,
        "Davon korrekt abgelehnt (F-13)": n_korrekt,
        "Quote %": round(100 * n_korrekt / n_neg, 1) if n_neg else None,
        "Vermutlich halluziniert (Testfall-Nr.)": ", ".join(str(n) for n in halluziniert_nr) or "—",
    }])

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_results.to_excel(writer, sheet_name="Evaluation Results", index=False)
        df_summary.to_excel(writer, sheet_name="Summary", index=False)
        df_overview.to_excel(writer, sheet_name="Metric Overview", index=False)
        df_issues.to_excel(writer, sheet_name="Issues", index=False)
        df_ziele.to_excel(writer, sheet_name="Beantwortbarkeit", index=False, startrow=1)
        df_ablehnung.to_excel(writer, sheet_name="Beantwortbarkeit", index=False,
                               startrow=len(df_ziele) + 5)
        wb = writer.book

        ws1 = writer.sheets["Evaluation Results"]
        _style_header(ws1, len(df_results.columns))
        _autosize(ws1, df_results)
        ws1.freeze_panes = "A2"

        ws2 = writer.sheets["Summary"]
        _style_header(ws2, len(df_summary.columns))
        _autosize(ws2, df_summary)

        ws3 = writer.sheets["Metric Overview"]
        _style_header(ws3, len(df_overview.columns))
        _autosize(ws3, df_overview)

        ws4 = writer.sheets["Issues"]
        _style_header(ws4, len(df_issues.columns) if len(df_issues.columns) else 4)
        if len(df_issues):
            _autosize(ws4, df_issues)
            for row_idx in range(2, len(df_issues) + 2):
                for col_idx in range(1, len(df_issues.columns) + 1):
                    ws4.cell(row=row_idx, column=col_idx).fill = ISSUE_FILL

        ws5 = writer.sheets["Beantwortbarkeit"]
        ws5["A1"] = f"Beantwortbare Testfaelle (n={len(beantwortbar_rows)}) — gegen Architektur-Zielwerte Kap. 4.7.2"
        ws5["A1"].font = Font(bold=True)
        _style_header(ws5, len(df_ziele.columns), row=2)
        _autosize(ws5, df_ziele)

        titel_zeile = len(df_ziele) + 4
        ws5.cell(row=titel_zeile, column=1,
                 value="Nicht-beantwortbare Testfaelle — F-13/UC-03 A2: korrekte Ablehnung statt Halluzination")
        ws5.cell(row=titel_zeile, column=1).font = Font(bold=True)
        _style_header(ws5, len(df_ablehnung.columns), row=len(df_ziele) + 6)

        # Balkendiagramm der Mittelwerte auf "Metric Overview"
        chart = BarChart()
        chart.title = "Durchschnittliche RAGAs-Metrikwerte"
        chart.y_axis.title = "Score (0-1)"
        chart.x_axis.title = "Metrik"
        chart.y_axis.scaling.min = 0
        chart.y_axis.scaling.max = 1
        data = Reference(ws3, min_col=2, max_col=2, min_row=1, max_row=len(df_overview) + 1)
        cats = Reference(ws3, min_col=1, max_col=1, min_row=2, max_row=len(df_overview) + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.width = 18
        chart.height = 10
        ws3.add_chart(chart, f"A{len(df_overview) + 4}")

        # Laufmetadaten oberhalb der Tabelle in "Summary" einfügen
        if run_meta:
            ws2.insert_rows(1, amount=len(run_meta) + 2)
            ws2["A1"] = "Lauf-Metadaten"
            ws2["A1"].font = Font(bold=True)
            r = 2
            for k, v in run_meta.items():
                ws2.cell(row=r, column=1, value=k)
                ws2.cell(row=r, column=2, value=str(v))
                r += 1
            _style_header(ws2, len(df_summary.columns), row=len(run_meta) + 2)

    return output_path
